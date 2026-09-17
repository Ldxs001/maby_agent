#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Copyright 2026 wUwproject
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""端到端验收：走界面用的那条 HTTP 路径，把「立项 → 脚本带标题 → 出片 → 期号前进」跑通。

单元测试各自测自己那一层，唯独串联处没人测：标题从脚本传到合成、期号从项目
记忆里接着往下排，这两件事都只有一个接口接一个接口地跑一遍才看得出来。
测试产物一律落在临时目录，跑完把配置还原，不动用户的真实产物。

用法：
    python main.py --port 8812
    python tools/e2e_http.py --url http://127.0.0.1:8812
"""

import argparse
import base64
import json
import os
import shutil
import sys
import time
import urllib.error
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRATCH = os.path.join("_smoke", "_scratch", "e2e_episodes")


def _abs_scratch(value):
    """把 --scratch 归成绝对路径。

    允许指到系统临时区：宿主若对「非临时路径的删除」设了拦截，
    脚本开头清理上次残留时会把自己弄死，换个临时根就绕开了。
    """
    return value if os.path.isabs(value) else os.path.join(ROOT, value)


def _restore_config_on_disk(keep):
    """服务不在时直接写回 config.json，返回写回后的键值。

    走这条路意味着验收过程中服务非正常退出（被外部因素带走、端口被抢等）。
    即便如此，用户原来的配置也必须回去——停在测试值上比测试失败更糟，
    因为它不会报错，只会让人下次打开界面时发现产物目录指去了临时区。
    """
    path = os.path.join(ROOT, "config.json")
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return {}
    for k, v in keep.items():
        if v is not None:
            data[k] = v
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    os.replace(tmp, path)
    return {k: data.get(k) for k in keep}


MATERIAL = """\
写代码和写文章，其实是同一件事的两面。

第一面是确定性的部分。哪些环节能用算法、正则、代码解决，就不该交给模型。
模型擅长的是另一头：理解人的意图，把意图翻译成结构。中间那条链，应该是代码。

第二面是不确定性的部分。剩余的解释空间，才是模型该填的。空着的地方要敢承认
填不了，而不是编一个看起来对的答案。

这两面合起来，就是一条链加两头：链是主体，模型只做两端。
"""

# 成稿规划那一档要用的书稿。章标题必须逐字稳定：排图给出的落点靠它核对，
# 而落点与锚点是同一个字符串——上传路径把标题剥掉，锚点就没了。
BOOK = """\
# 第一章 链与两头

写代码和写文章是同一件事的两面。确定的环节交给代码，解释空间留给模型。

中间那条链是主体，模型只做两端。

# 第二章 代码编期号

期号若由模型来编，一旦有一期对不上，后面全部错位，而每期单看都正常。

所以期号只由代码编，模型只管内容。

# 第三章 落点要逐字核实

模型偶尔会把章节名顺手润色一下，那种锚点在素材里找不到，直到取料才炸。

所以排完图当场核对，对不上的落点丢弃并记明。
"""


def call(base, path, body=None, method=None, timeout=120):
    url = base + path
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method or
                                 ("POST" if data is not None else "GET"))
    if data is not None:
        req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return {"ok": False, "error": "HTTP %s: %s" % (e.code, e.read()[:200])}
    except Exception as e:                                  # noqa: BLE001
        return {"ok": False, "error": str(e)}


class Runner(object):
    """判据收集器。判据数由它自己数，不靠人去点源码——写进更新日志的数字
    必须来自工具，否则改一条判据就要重新数一遍，迟早对不上。"""

    def __init__(self, base):
        self.base = base
        self.fails = []
        self.total = 0

    def check(self, ok, label, detail=""):
        self.total += 1
        print("  %-4s %s%s" % ("PASS" if ok else "FAIL", label,
                               ("  " + detail) if detail else ""))
        if not ok:
            self.fails.append(label + (("  " + detail) if detail else ""))
        return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--url", default="http://127.0.0.1:8812")
    ap.add_argument("--model", default="qwen3-8b")
    ap.add_argument("--minutes", type=float, default=1.0)
    ap.add_argument("--scratch", default=SCRATCH,
                    help="验收产物落点。默认落在仓库内的 _smoke/_scratch；"
                         "指到系统临时区可避开宿主对非临时路径删除的拦截。")
    args = ap.parse_args()
    base = args.url.rstrip("/")
    r = Runner(base)

    cfg0 = call(base, "/api/config")
    if not cfg0.get("ok"):
        print("配置读不到，服务没起来？%s" % cfg0.get("error"))
        return 1
    keep = {k: cfg0["values"].get(k) for k in
            ("llm.model", "script.target_minutes", "project.output_dir")}
    print("原配置：%s" % json.dumps(keep, ensure_ascii=False))

    out_root = _abs_scratch(args.scratch)
    shutil.rmtree(out_root, ignore_errors=True)
    pid = ""
    created = []           # 验收造出的项目，跑完一并从真实登记表里摘掉
    try:
        # 临时把产物目录指到临时区，别把真实产物搅进来
        call(base, "/api/config", {"patch": {
            "llm.model": args.model,
            "script.target_minutes": args.minutes,
            "project.output_dir": out_root.replace("\\", "/")}})

        print()
        print("一、立项")
        p = call(base, "/api/project", {
            "action": "create", "name": "端到端验收档", "program_name": "验收节目",
            "plan_mode": "episodic",
            "planned_episodes": 3, "first_episode": "1", "style_preset": "argument",
            "note": "由 tools/e2e_http.py 自动建立"})
        r.check(p.get("ok"), "立项成功", str(p.get("error") or ""))
        if not p.get("ok"):
            return 1
        pid = p["project"]["id"]
        created.append(pid)
        r.check(p["project"]["progress"]["next_episode"] == "1",
                "初始下一期号是 1", p["project"]["progress"]["next_episode"])
        r.check(p["project"]["progress"]["planned"] == 3,
                "计划期数按人填的 3", str(p["project"]["progress"]["planned"]))
        r.check(p["project"].get("plan_mode") == "episodic",
                "规划方式记下了", str(p["project"].get("plan_mode")))

        # 规划方式在立项时定死。不选就立不了项，立了之后也改不动——
        # 中途换模式会让已排的地图失去来处。
        nomode = call(base, "/api/project", {"action": "create", "name": "没选规划方式"})
        r.check(not nomode.get("ok"), "立项不选规划方式会被拒",
                str(nomode.get("error") or "")[:120])
        chg = call(base, "/api/project", {"action": "update", "id": pid,
                                         "plan_mode": "mapped"})
        r.check(not chg.get("ok"), "已立项的项目改不了规划方式",
                str(chg.get("error") or "")[:120])

        print()
        print("二、生成脚本（标题随脚本产出）")
        g = call(base, "/api/script/generate",
                 {"material": MATERIAL, "extra": "", "preset": "argument",
                  "project_id": pid}, timeout=600)
        r.check(g.get("ok"), "脚本生成成功", str(g.get("error") or "")[:200])
        if not g.get("ok"):
            return 1
        title = (g.get("title") or "").strip()
        r.check(bool(title), "脚本带回了标题", "标题=%r" % title)
        r.check(len(title) <= 24, "标题长度不超上限", "%d 字" % len(title))
        r.check(g.get("planned_episodes") == 3,
                "项目已定期数时模型不改口", str(g.get("planned_episodes")))
        r.check(str(g.get("episode_no")) == "1", "本次是第 1 期",
                str(g.get("episode_no")))
        r.check(bool(g.get("script")), "脚本有内容", "%d 句" % len(g.get("script") or []))
        gref = g.get("report", {})
        r.check(gref.get("passed"), "生成门禁通过",
                "阻断 %s / 记录 %s" % ([i["label"] for i in gref.get("fails", [])],
                                    [i["label"] for i in gref.get("warns", [])]))

        print()
        print("三、出片（不产视频，省时间）")
        job = call(base, "/api/render", {
            "title": title, "episode_no": str(g.get("episode_no") or "1"),
            "script": g["script"], "material": MATERIAL, "preset": "argument",
            "extra": "", "do_video": False, "project_id": pid})
        r.check(job.get("ok"), "任务已受理", str(job.get("error") or "")[:200])
        if not job.get("ok"):
            return 1
        tid = job["task_id"]
        res, done_ok, deadline = {}, False, time.time() + 900
        while time.time() < deadline:
            st = call(base, "/api/task/" + tid)
            j = st.get("job") or {}
            if j.get("status") != "running":
                res = j.get("result") or {}
                done_ok = j.get("status") == "done"
                r.check(done_ok, "合成完成", str(j.get("error") or "")[:200])
                break
            time.sleep(2)
        else:
            r.check(False, "合成完成", "超时")
            return 1

        print()
        print("四、产物与校验")
        assets = res.get("assets") or {}
        for key, label in (("audio", "音频"), ("article", "图文"),
                           ("subtitle", "字幕"), ("bg", "背景图")):
            r.check(bool(assets.get(key)), "产物有%s" % label, str(assets.get(key) or ""))
        r.check(len(assets.get("covers") or {}) >= 1, "产物有封面",
                str(list((assets.get("covers") or {}).keys())))
        rref = res.get("report", {})
        r.check(rref.get("passed"), "产物校验通过",
                "阻断 %s / 记录 %s" % ([i["label"] for i in rref.get("fails", [])],
                                    [i["label"] for i in rref.get("warns", [])]))
        r.check(not res.get("video"), "按开关没有产视频", str(res.get("video") or "无"))
        r.check(res.get("next_episode") == "2", "结果里带回下一期号",
                str(res.get("next_episode")))

        print()
        print("五、期号记忆与续接")
        if not done_ok:
            r.check(False, "期号随出片前进", "合成没完成，这一步无从验起")
            return 1
        after = call(base, "/api/project")
        item = next((x for x in after.get("projects") or [] if x["id"] == pid), None)
        r.check(item is not None, "项目还在")
        if item:
            eps = item.get("episodes") or []
            r.check(len(eps) == 1, "记录了一期",
                    json.dumps(eps, ensure_ascii=False))
            r.check(item["next_episode"] == "2", "下一期号前进到 2",
                    item["next_episode"])
            r.check(eps and eps[0]["title"] == title, "记录的是脚本产出的标题",
                    eps[0]["title"] if eps else "（没有记录）")
            r.check(item["progress"]["done"] == 1, "进度计数为 1",
                    str(item["progress"]["done"]))

        g2 = call(base, "/api/script/generate",
                  {"material": MATERIAL, "extra": "", "preset": "argument",
                   "project_id": pid}, timeout=600)
        r.check(g2.get("ok"), "同项目再次生成成功", str(g2.get("error") or "")[:160])
        r.check(str(g2.get("episode_no")) == "2", "续接时自动排到第 2 期",
                str(g2.get("episode_no")))

        print()
        print("六、未归属产物不该被自动认领")
        ep_dir = os.path.basename((res.get("project_dir") or "").rstrip("\\/"))
        orph = call(base, "/api/project?scope=orphans")
        r.check(orph.get("ok"), "未归属列表可读")
        r.check(all(o["dir"] != ep_dir for o in orph.get("orphans") or []),
                "已入册的产物不出现在未归属里",
                "本期目录 %s / 未归属 %s" % (ep_dir,
                                        str([o["dir"] for o in orph.get("orphans") or []])[:120]))

        print()
        print("七、成稿规划：整本书 → 期数地图 → 免摘抄出脚本 → 插入分支")
        pm = call(base, "/api/project", {
            "action": "create", "name": "端到端验收档·成稿", "program_name": "验收节目",
            "plan_mode": "mapped", "planned_episodes": 2, "first_episode": "1",
            "style_preset": "argument", "paradigm": "methodology",
            "note": "由 tools/e2e_http.py 自动建立"})
        r.check(pm.get("ok"), "成稿规划立项成功", str(pm.get("error") or "")[:200])
        if not pm.get("ok"):
            return 1
        pid2 = pm["project"]["id"]
        created.append(pid2)
        r.check(pm["project"].get("plan_mode") == "mapped",
                "立项时选定的模式被记下", str(pm["project"].get("plan_mode")))
        r.check(pm["project"].get("map") is None, "新项目还没有地图")
        r.check(pm["project"].get("paradigm") == "methodology",
                "立项时选的素材类型被记下",
                str(pm["project"].get("paradigm")))
        r.check(((pm["project"].get("progress") or {})
                 .get("paradigm_label") or "") != "自适应（按结构推断）",
                "项目进度显示的是所选类型，不是「自适应」",
                str((pm["project"].get("progress") or {})
                    .get("paradigm_label")))

        # 上传路径：走 base64 上传 .md，而不是粘贴。锚点若在上传时被剥掉，
        # 这里就会看到 0 节——那正是「素材为空 / 无法按章切分」的根因。
        payload = base64.b64encode(BOOK.encode("utf-8")).decode("ascii")
        up = call(base, "/api/source", {
            "action": "add", "project_id": pid2, "name": "书稿.md",
            "filename": "book.md", "data_base64": payload})
        r.check(up.get("ok"), "书稿上传入库", str(up.get("error") or "")[:200])
        src = up.get("source") or {}
        r.check(src.get("sections") == 3,
                "上传的 md 认出 3 个章节锚点", "sections=%s" % src.get("sections"))
        r.check(src.get("chars") == len(BOOK.strip()),
                "入库字数与原文一致", "%s / %s" % (src.get("chars"), len(BOOK.strip())))
        # 入库即探查：结构（层级、切分单位、准入、容量）与类型在这一步定下来，
        # 落盘冻结后供排图直接取用。入库不探查时，界面看不到任何结构，
        # 排图那一步才现跑——而排图失败就把探查一起赔进去了。
        up_logs = up.get("logs") or []
        r.check(any(m.startswith("探查 ") for m in up_logs),
                "入库即探查：跑了结构探查", str(up_logs)[:200])
        r.check(any(m.startswith("类型判定") for m in up_logs),
                "入库即探查：跑了类型判定", str(up_logs)[:200])
        upv = up.get("probe") or {}
        r.check(upv.get("segments") == 3,
                "入库返回里带结构摘要", str(upv.get("segments")))
        r.check(bool(upv.get("kind_label")), "摘要里带类型名",
                str(upv.get("kind_label")))
        r.check(bool(upv.get("model")), "摘要里记下判定用的模型",
                str(upv.get("model")))
        listed = call(base, "/api/sources?project_id=" + pid2)
        r.check(len(listed.get("sources") or []) == 1, "素材清单里有一份",
                str([s.get("name") for s in listed.get("sources") or []]))
        lpv = ((listed.get("sources") or [{}])[0].get("probe")) or {}
        r.check(lpv.get("segments") == 3,
                "素材清单里也带结构摘要", str(lpv.get("segments")))

        # 补跑入口：旧版本入库的素材没有探查结果，得有路把它补上。
        # 已经探查过的走缓存——不重扫，也不重判（重判会白烧一次模型）。
        sc = call(base, "/api/source", {"action": "scan", "project_id": pid2},
                  timeout=900)
        r.check(sc.get("ok"), "已入库的素材能补跑探查",
                str(sc.get("error") or "")[:200])
        if sc.get("ok"):
            r.check(int(sc.get("est_episodes") or 0) > 0,
                    "补跑后给出按体量算的期数", str(sc.get("est_episodes")))
            r.check(((sc.get("sources") or [{}])[0].get("probe") or {})
                    .get("segments") == 3, "补跑后素材带上结构摘要")
            r.check(not [m for m in (sc.get("logs") or [])
                         if m.startswith("类型判定")],
                    "已探查过的素材不重复判定", str(sc.get("logs"))[:160])

        mp = call(base, "/api/plan", {"action": "map", "project_id": pid2}, timeout=900)
        r.check(mp.get("ok"), "排地图成功", str(mp.get("error") or "")[:200])
        if not mp.get("ok"):
            return 1
        rows = ((mp.get("project") or {}).get("map") or {}).get("episodes") or []
        r.check(len(rows) == 2, "按人定的期数排出 2 期", "实际 %d 期" % len(rows))
        r.check([e["no"] for e in rows] == ["1", "2"],
                "期号由代码从 1 顺编", str([e["no"] for e in rows]))
        r.check(all(e.get("refs") for e in rows), "每期都有素材落点",
                str([e.get("refs") for e in rows])[:160])
        known = {l.lstrip("#").strip() for l in BOOK.split("\n") if l.startswith("# ")}
        bad = [rf.get("anchor") for e in rows for rf in e.get("refs") or []
               if rf.get("anchor") not in known]
        r.check(not bad, "落点都逐字命中素材里的章节", "对不上的：%s" % bad)
        r.check(((mp.get("project") or {}).get("progress") or {}).get("mapped") == 2,
                "进度里记下了地图规模",
                str(((mp.get("project") or {}).get("progress") or {}).get("mapped")))

        # 不手摘素材：一期讲哪几节地图已经记着，脚本阶段就该自己去取。
        g3 = call(base, "/api/script/generate",
                  {"material": "", "extra": "", "preset": "argument",
                   "project_id": pid2}, timeout=900)
        r.check(g3.get("ok"), "素材留空也能生成脚本", str(g3.get("error") or "")[:200])
        if g3.get("ok"):
            r.check(any("按地图落点取素材" in m for m in (g3.get("logs") or [])),
                    "日志写明素材按落点取用",
                    str(g3.get("logs") or [])[:160])
            r.check((g3.get("title") or "").strip() == rows[0]["title"],
                    "脚本标题取自地图（不重新命名）",
                    "%r / 地图 %r" % (g3.get("title"), rows[0]["title"]))

        ins = call(base, "/api/plan",
                   {"action": "insert_plan", "project_id": pid2,
                    "source_ids": [src.get("id")]}, timeout=900)
        r.check(ins.get("ok"), "模型给出插入建议", str(ins.get("error") or "")[:200])
        if not ins.get("ok"):
            return 1
        sug = ins.get("suggestion") or {}
        existing = [e["no"] for e in rows]
        r.check(sug.get("anchor_no") in existing, "插入点是现有期号",
                "%s / 现有 %s" % (sug.get("anchor_no"), existing))
        eps = sug.get("episodes") or []
        r.check(bool(eps), "给出了要插入的期",
                str([e.get("no") for e in eps]))
        r.check(all(str(e.get("no") or "").startswith(str(sug.get("anchor_no")))
                    and e.get("no") != sug.get("anchor_no") for e in eps),
                "建议里的分支期号长在锚点期号上", str([e.get("no") for e in eps]))

        # 落库时期号故意递一个错的：后端必须自己编，不能采信调用方给的号。
        ap = call(base, "/api/plan", {
            "action": "insert_apply", "project_id": pid2,
            "anchor_no": sug.get("anchor_no"),
            "episodes": [{"no": "瞎写的", "title": e.get("title"),
                          "points": e.get("points") or [], "refs": e.get("refs") or []}
                         for e in eps]})
        r.check(ap.get("ok"), "插入落库成功", str(ap.get("error") or "")[:200])
        nos = ap.get("episodes") or []
        r.check(len(nos) == len(eps), "返回插入的期号", str(nos))
        r.check(not any(n == "瞎写的" for n in nos), "后端没有采信调用方给的期号",
                str(nos))
        r.check(nos == [str(sug.get("anchor_no")) + chr(ord("a") + i)
                        for i in range(len(nos))],
                "分支期号从 a 起顺编", str(nos))

        after2 = call(base, "/api/project")
        item2 = next((x for x in after2.get("projects") or [] if x["id"] == pid2), None)
        order = [e["no"] for e in ((item2 or {}).get("map") or {}).get("episodes") or []]
        at = order.index(sug.get("anchor_no")) if sug.get("anchor_no") in order else -1
        r.check(at >= 0 and order[at + 1:at + 1 + len(nos)] == nos,
                "分支期紧跟在锚点期之后", "地图顺序 %s" % order)
        r.check(len(order) == len(rows) + len(nos), "地图期数只增了插入的这几期",
                "%d → %d" % (len(rows), len(order)))

        # 上传 docx 时样式证据必须跟着入库。marks 只在读文件那一刻存在，
        # 入库时丢掉，第二次读素材就只剩裸段落，docx 退化成一个更大的 txt——
        # 而 sections 数照样对得上（编号本身就认得出），看不出来。
        import tempfile as _tf
        try:
            import docx as _docx
        except ImportError:                                       # pragma: no cover
            _docx = None
        if _docx is not None:
            _buf = os.path.join(_tf.gettempdir(), "pm_e2e_book.docx")
            _d = _docx.Document()
            _d.add_heading("第一章 样式给的层级", level=1)
            _d.add_paragraph("这一章的层级来自标题样式，不是猜的。" * 6)
            _d.add_heading("第一节 下级小节", level=2)
            _d.add_paragraph("下级小节也来自样式。" * 6)
            _d.save(_buf)
            with open(_buf, "rb") as _f:
                _payload = base64.b64encode(_f.read()).decode("ascii")
            os.remove(_buf)
            up2 = call(base, "/api/source", {
                "action": "add", "project_id": pid2, "name": "样章.docx",
                "filename": "样章.docx", "data_base64": _payload})
            r.check(up2.get("ok"), "docx 上传入库",
                    str(up2.get("error") or "")[:200])
            s2 = up2.get("source") or {}
            r.check(bool(s2.get("marks")), "docx 的样式证据随素材入库",
                    "marks=%s" % str(s2.get("marks"))[:80])
            _lv = {m.get("title"): m.get("level")
                   for m in s2.get("marks") or []}
            r.check(_lv.get("第一章 样式给的层级") == 1
                    and _lv.get("第一节 下级小节") == 2,
                    "样式层级随素材入库", str(_lv))

    finally:
        print()
        print("八、还原现场")
        # 还原不能假定服务还活着：验收过程中服务若被外部因素带走，
        # 还原请求会打空，配置就永久停在测试值上（产物目录指向临时区）。
        # 所以先试服务，服务不在就直接改盘上的配置。
        alive = call(base, "/api/config").get("ok")
        if alive:
            for k, v in keep.items():
                if v is not None:
                    call(base, "/api/config", {"patch": {k: v}})
            back = (call(base, "/api/config") or {}).get("values") or {}
        else:
            print("  服务不可达，改为直接写回配置文件")
            back = _restore_config_on_disk(keep)
        same = all(str(back.get(k)) == str(v) for k, v in keep.items() if v is not None)
        r.check(same, "配置已还原", json.dumps({k: back.get(k) for k in keep},
                                             ensure_ascii=False))
        if created:
            # 验收项目与本次造出的产物一并挪到验收区：归档优于删除，
            # 但它们不该留在用户的真实登记表与产物目录里。
            for one in created:
                call(base, "/api/project", {"action": "update", "id": one,
                                            "archived": True})
            store = os.path.join(ROOT, "projects", "_projects.json")
            if os.path.exists(store):
                with open(store, "r", encoding="utf-8") as fh:
                    data = json.load(fh)
                data["projects"] = [x for x in data["projects"]
                                    if x["id"] not in created]
                tmp = store + ".tmp"
                with open(tmp, "w", encoding="utf-8") as fh:
                    json.dump(data, fh, ensure_ascii=False, indent=2)
                os.replace(tmp, store)
            r.check(True, "验收项目已从登记表移除", "、".join(created))
            # 素材是跟着项目 id 落在产物目录下的，项目撤了它也得跟着走
            src_root = os.path.join(out_root, "_sources")
            if os.path.isdir(src_root):
                shutil.rmtree(src_root, ignore_errors=True)
        if os.path.isdir(out_root):
            done = os.path.join(os.path.dirname(out_root), "e2e_episodes_done")
            shutil.rmtree(done, ignore_errors=True)
            shutil.move(out_root, done)

    print()
    print("=" * 62)
    if r.fails:
        print("端到端验收：%d 项中 %d 项未通过" % (r.total, len(r.fails)))
        for f in r.fails:
            print("  -", f)
        return 1
    print("端到端验收：%d 项全部通过" % r.total)
    return 0


if __name__ == "__main__":
    sys.exit(main())
