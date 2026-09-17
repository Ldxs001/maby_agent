#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Copyright 2026 wUwproject
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""结构探查：L1 确定性、L2 可复现、准入前移。

探查错了不会当场报错，它会一路传到切章、排图、取料，最后表现为「地图排得
乱七八糟」。所以这里的判据全是**不变量**而非「像不像」：同输入同输出、编号
认得全、围栏内不算标题、空正文进不了图、杂项默认不播。
"""

import json
import os
import shutil
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from podcast_maker import ingest                                          # noqa: E402
from podcast_maker import llm_client as llm_mod                           # noqa: E402
from podcast_maker import paradigms as PG                                 # noqa: E402
from podcast_maker import probe                                           # noqa: E402
from podcast_maker import project_store as P                              # noqa: E402
from podcast_maker import source_store as S                               # noqa: E402

CFG = {"script.target_minutes": 25.0, "script.map_head_chars": 280,
       "script.map_max_episodes": 60}

BOOK = """# 第一章 起点

这一章讲事情的由来。

# 第二章 转折

这一章讲分岔路口。

# 第三章 收束

这一章讲如何收尾。
"""


def body(text, n=6):
    return "正文内容。" * 30


class TestNumbering(unittest.TestCase):
    """C 路是纯文本唯一的依靠：编号写在正文里，抽成纯文本也还在。

    这些形态用户点名要照顾到——`(一)`、`一、`、`1.`、`1.1`、`a)`、`第X章`。
    少认一种，那种稿子就会被当成「没有结构」，整本排成一期。
    """

    def titles(self, text):
        ir = probe.scan(text)
        return [s["title"] for s in ir["segments"]]

    def test_cn_paren_number(self):
        got = self.titles("（一）背景\n%s\n（二）方法\n%s" % (body(""), body("")))
        self.assertIn("（一）背景", got)
        self.assertIn("（二）方法", got)

    def test_half_paren_number(self):
        got = self.titles("(一) 背景\n%s\n(二) 方法\n%s" % (body(""), body("")))
        self.assertIn("(一) 背景", got)

    def test_cn_dun_number(self):
        got = self.titles("一、背景\n%s\n二、方法\n%s" % (body(""), body("")))
        self.assertIn("一、背景", got)

    def test_dotted_number(self):
        got = self.titles("1. 背景\n%s\n2. 方法\n%s" % (body(""), body("")))
        self.assertIn("1. 背景", got)

    def test_dotted_sub_number(self):
        got = self.titles("1.1 观察\n%s\n1.2 判断\n%s" % (body(""), body("")))
        self.assertIn("1.1 观察", got)
        self.assertIn("1.2 判断", got)

    def test_letter_paren(self):
        got = self.titles("a) 起点\n%s\nb) 终点\n%s" % (body(""), body("")))
        self.assertIn("a) 起点", got)

    def test_cn_chapter(self):
        got = self.titles("第一章 起点\n%s\n第二章 转折\n%s" % (body(""), body("")))
        self.assertIn("第一章 起点", got)
        self.assertIn("第二章 转折", got)

    def test_en_chapter(self):
        got = self.titles("Chapter 1 Origin\n%s\nChapter 2 Turn\n%s"
                          % (body(""), body("")))
        self.assertIn("Chapter 1 Origin", got)

    def test_roman(self):
        got = self.titles("IV. 第四部\n%s\nV. 第五部\n%s" % (body(""), body("")))
        self.assertIn("IV. 第四部", got)

    def test_branch_letter(self):
        # `08a 探测的边界`：同一数字前缀带不同字母，是同一篇的续篇
        got = self.titles("08 协作\n%s\n08a 探测的边界\n%s" % (body(""), body("")))
        self.assertIn("08a 探测的边界", got)

    def test_bare_text_is_not_invented(self):
        # 没有标记、没有样式、正文里也没有编号 → 不猜，明说未识别
        ir = probe.scan(body("", 20))
        self.assertEqual(len(ir["segments"]), 1)
        self.assertEqual(ir["segments"][0]["title"], "全文")
        self.assertTrue(any("未识别" in w for w in ir["warnings"]))


class TestFenceAndNoise(unittest.TestCase):
    """围栏内的 `#` 是代码注释，不是章节标题。

    从前只按 `#{1,6}` 判标题，于是书稿里用 `#` 写的代码行被判成了章——
    它把正文切走一块，还让那一篇的体量算少了，排期跟着错。
    """

    def titles(self, text):
        return [s["title"] for s in probe.scan(text)["segments"]]

    def test_hash_inside_fence_is_not_a_title(self):
        text = ("# 真标题\n\n正文。\n\n```text\n# 这是配置注释\n再兜底\n```\n\n"
                "更多正文。")
        got = self.titles(text)
        self.assertIn("真标题", got)
        self.assertNotIn("# 这是配置注释", got)

    def test_tilde_fence(self):
        got = self.titles("# 真标题\n\n~~~\n# 注释\n~~~\n\n正文。")
        self.assertNotIn("# 注释", got)

    def test_table_row_is_not_a_title(self):
        got = self.titles("# 真标题\n\n| 项 | 值 |\n|---|---|\n| a | b |\n")
        self.assertEqual(got, ["真标题"])

    def test_body_sentence_with_number_is_not_a_title(self):
        # 「1. 这一点很重要。」是一句话，不是标题——句末标点加长度否决
        got = self.titles("# 真标题\n\n1. 这一点很重要，必须先说清楚再做别的事情。\n")
        self.assertEqual(got, ["真标题"])


class TestSeries(unittest.TestCase):
    """同系列该合并排期，这是整合依据里最硬的一条，不能靠模型猜。"""

    def test_branch_family(self):
        text = "".join("# 08%s 第%s节\n\n%s\n\n" % (c, c, body(""))
                       for c in ("a", "b", "c"))
        ir = probe.scan(text)
        keys = [g["key"] for g in ir["series"]]
        self.assertTrue(any(g["titles"] and len(g["titles"]) == 3
                            for g in ir["series"]), keys)

    def test_text_prefix_family(self):
        text = ("# 重型技能构建的方法论沉淀：变量治理\n\n%s\n\n"
                "# 重型技能构建的方法论沉淀：路径治理\n\n%s\n" % (body(""), body("")))
        ir = probe.scan(text)
        self.assertTrue(ir["series"], "同前缀的两篇没有被归成一组")

    def test_numbered_prefix_is_not_a_family(self):
        # `1.2 判断` 的 `1` 是它自己的编号，不是系列名
        text = "# 1.2 判断\n\n%s\n\n# 2.2 判断\n\n%s\n" % (body(""), body(""))
        for g in probe.scan(text)["series"]:
            self.assertNotIn(g["key"], ("前缀 1", "前缀 2"))


class TestMiscAndAdmissible(unittest.TestCase):
    def test_appendix_is_not_played_by_default(self):
        ir = probe.scan("# 正文一\n\n%s\n\n# 附录 A\n\n%s\n" % (body(""), body("")))
        by = {s["title"]: s for s in ir["segments"]}
        self.assertEqual(by["附录 A"]["misc"], "drop")

    def test_preface_is_played(self):
        ir = probe.scan("# 序言\n\n%s\n\n# 正文一\n\n%s\n" % (body(""), body("")))
        by = {s["title"]: s for s in ir["segments"]}
        self.assertEqual(by["序言"]["misc"], "keep")

    def test_empty_body_is_not_admissible(self):
        # 标题挨着标题、中间没有正文 → 切不出料，排图前就要挡掉
        ir = probe.scan("# 甲\n# 乙\n\n%s\n" % body(""))
        ok, bad = probe.admissible(ir)
        self.assertTrue(bad, "空正文的条目没有被挡下")
        self.assertTrue(all(s["title"] != "甲" for s in ok))

    def test_units_picks_unit_level(self):
        text = ("# 第一篇\n\n%s\n\n## 第一章\n\n%s\n\n## 第二章\n\n%s\n"
                % (body(""), body(""), body("")))
        ir = probe.scan(text)
        ir["unit_level"] = 1
        got = [s["title"] for s in probe.units(ir)]
        self.assertIn("第一篇", got)


class TestMiscFollowsTheCard(unittest.TestCase):
    """附属页判定要认类型：**卡上写的优先，卡上没列的才落通用兜底**。

    接线前这一步只认字面（通用名单），卡上那份清单没人读。两处结论不同的地方
    真实存在——论文的「致谢」在卡上是不播的，通用名单却把它归进保留（散文、
    纪实里的致谢确实是作者的话）。只认字面，这类分歧就永远按通用名单判。
    """

    TEXT = ("# 正文一\n\n%s\n\n# 致谢\n\n%s\n\n# 参考答案\n\n%s\n\n# 摘要\n\n%s\n"
            % (body(""), body(""), body(""), body("")))

    def misc_of(self, card=None):
        ir = probe.scan(self.TEXT, None, probe.HEAD_CHARS, card=card)
        return {s["title"]: s["misc"] for s in ir["segments"]}

    def test_no_card_falls_back_to_the_shared_list(self):
        by = self.misc_of(None)
        self.assertEqual(by["致谢"], "keep", "没卡时按通用名单：致谢是作者的话")
        self.assertEqual(by["参考答案"], "", "通用名单里没有这一条，落回正文")
        self.assertEqual(by["摘要"], "", "同上")

    def test_card_beats_the_shared_list(self):
        by = self.misc_of(PG.get("paper"))
        self.assertEqual(by["致谢"], "drop", "论文的致谢是感谢导师与基金，不播")
        self.assertEqual(by["摘要"], "keep", "摘要该播，且通用名单里没有它")

    def test_card_does_not_touch_what_it_does_not_list(self):
        by = self.misc_of(PG.get("methodology"))
        self.assertEqual(by["参考答案"], "", "方法论卡没写它，仍落通用兜底")

    def test_only_the_listing_card_changes_the_verdict(self):
        base = self.misc_of(None)
        tut = self.misc_of(PG.get("tutorial"))
        self.assertNotEqual(base["参考答案"], tut["参考答案"],
                            "教程卡该把「参考答案」拦下")


class TestOutlier(unittest.TestCase):
    def test_unmarked_sibling_is_flagged(self):
        # 同层标题大多带编号，冒出来一条不带编号又很短的，多半是误标
        # 真场景：三层结构里，最后冒出一条浅层、无编号、又很短的标题
        text = "".join("# %d. 第%d章\n\n%s\n\n## %d.1 小节\n\n%s\n\n"
                       "### %d.1.1 细目\n\n%s\n\n"
                       % (i, i, body(""), i, body(""), i, body(""))
                       for i in range(1, 5))
        text += "# 再兜底（后置，幂等覆盖）\n\n短。\n"
        ir = probe.scan(text)
        flagged = [s["title"] for s in ir["segments"] if s.get("warn")]
        self.assertIn("再兜底（后置，幂等覆盖）", flagged)


class TestCapacity(unittest.TestCase):
    """期数是算出来的，不是问出来的。同一份素材两次必须同一个数。"""

    def test_deterministic(self):
        a = probe.capacity(CFG)
        b = probe.capacity(dict(CFG))
        self.assertEqual(a, b)

    def test_scales_with_target_minutes(self):
        short = probe.capacity(dict(CFG, **{"script.target_minutes": 10.0}))
        long_ = probe.capacity(dict(CFG, **{"script.target_minutes": 40.0}))
        self.assertLess(short, long_)

    def test_capacity_is_target_times_ratio(self):
        # 容量口径 = 成稿目标字数 × 压缩档。三把尺子（下钻、排图参照、
        # 期数折算）共用这一个数，分母（目标字数）与档位各自可验
        for ratio in probe.RATIOS:
            cfg = dict(CFG, **{"script.compress_ratio": ratio})
            self.assertEqual(probe.capacity(cfg),
                             round(probe.target_chars(cfg) * ratio))

    def test_ratio_of_clamps_to_known_gears(self):
        # 手改配置文件写出非法档位时回默认——宁可保守，不拿没对过账的倍数画图
        self.assertEqual(probe.ratio_of({"script.compress_ratio": 3}), 3)
        self.assertEqual(probe.ratio_of({"script.compress_ratio": 10}), 10)
        self.assertEqual(probe.ratio_of({"script.compress_ratio": 7}), 5)
        self.assertEqual(probe.ratio_of({"script.compress_ratio": "x"}), 5)
        self.assertEqual(probe.ratio_of({}), 5)

    def test_estimate_episodes(self):
        self.assertEqual(probe.estimate_episodes(12000, 6000), 2)
        self.assertEqual(probe.estimate_episodes(0, 6000), 1)
        self.assertEqual(probe.estimate_episodes(12000, 0), 1)

    def test_fill_capacity_writes_into_ir(self):
        ir = probe.scan("# 一\n\n%s\n" % body(""))
        probe.fill_capacity(ir, CFG)
        self.assertGreater(ir["capacity"], 0)
        self.assertGreaterEqual(ir["est_episodes"], 1)


class TestUnitPosition(unittest.TestCase):
    """位置与章节上下：排图靠它合并与切分，取料靠它对上号。

    位置不是装饰。凝缩压的是哪一块、地图上标的落点、写脚本时取的那一段，
    三处必须是同一个答案——两处口径一漂，地图看着完整，取的却是别的段落。
    """

    TEXT = ("# 第一部 总论\n\n概述。\n\n## 第一章 起点\n\n%s\n\n"
            "## 第二章 转折\n\n%s\n" % (body(""), body("")))

    def test_scan_records_end_and_path(self):
        ir = probe.scan(self.TEXT)
        segs = ir["segments"]
        self.assertTrue(all("end" in s for s in segs))
        h1 = [s for s in segs if s["level"] == 1]
        h2 = [s for s in segs if s["level"] == 2]
        self.assertEqual(h1[0]["path"], "")
        self.assertEqual([s["path"] for s in h2], ["第一部 总论"] * 2)
        self.assertEqual(probe.loc_text(h2[0]),
                         "第 %d–%d 行" % (h2[0]["line"] + 1, h2[0]["end"]))

    def test_scan_position_matches_what_slicing_takes(self):
        # 结构自算的位置与取料实际切走的那一段必须一致。两处口径漂了，地图上
        # 标的落点就是假证据——看着有位置，取的却是别的段落。
        ir = probe.scan(self.TEXT)
        for s in ir["segments"]:
            _got, meta = ingest.slice_by_anchor(self.TEXT, s["title"], [],
                                                line=s["line"])
            self.assertEqual(meta["line_start"], s["line"] + 1, s["title"])
            self.assertEqual(meta["line_end"], s["end"], s["title"])

    def test_structure_brief_gives_the_top_level_whole(self):
        # 最外一层是「篇」的候选，正是模型要在其中做选择的对象；抽样抽掉中间
        # 几条，正好把最该看的主体章节抽走了。更深的层才抽样。
        segs = ([{"level": 1, "title": "篇%02d" % i} for i in range(1, 41)]
                + [{"level": 2, "title": "节%03d" % i} for i in range(1, 201)])
        brief = probe._structure_brief(segs)
        self.assertIn("篇20", brief)
        self.assertIn("篇40", brief)
        self.assertIn("省略", brief)

    def test_unit_cost_leaves_out_appendix_and_empty(self):
        # 判层用的条数必须是「在播且切得出正文」那一份：把附录、空壳也算进去，
        # 那个区间量到的就不是「能排成多少期」。
        ir = probe.scan("# 第一章 A\n\n%s\n\n# 第二章 B\n\n%s\n\n"
                        "# 参考文献\n\n条目一。\n\n# 空壳\n"
                        % (body(""), body("")))
        self.assertEqual(ir["level_counts"]["1"], 4)
        self.assertEqual(probe._unit_cost(ir["segments"]).get(1), 2)

    def test_position_is_backfilled_from_what_slicing_took(self):
        s = {"title": "第一章 起点", "line": 9, "end": 48}
        drift = probe._mark_loc(
            s, {"line_start": 10, "line_end": 48, "chars": 500})
        self.assertEqual(drift, 0)
        self.assertEqual(s["loc"], {"start": 10, "end": 48, "chars": 500})
        self.assertEqual(probe.loc_text(s), "第 10–48 行")

    def test_position_drift_is_flagged(self):
        # 两处口径漂了就出声：位置一旦对不上，地图上的落点就成了假证据
        s = {"title": "第一章 起点", "line": 9, "end": 48}
        self.assertEqual(probe._mark_loc(
            s, {"line_start": 30, "line_end": 60}), 1)


class FakeLLM(object):
    def __init__(self, reply=None, boom=False):
        self.reply = reply
        self.boom = boom
        self.calls = []

    def chat(self, messages, temperature=0.8, max_tokens=8192, json_schema=None,
             **kw):
        rec = {"messages": messages, "temperature": temperature,
               "schema": json_schema}
        rec.update(kw)
        self.calls.append(rec)
        if self.boom:
            raise RuntimeError("后端不可用")
        return self.reply, {"model": "fake"}


class TestClassify(unittest.TestCase):
    """L2 是分类不是生成：输出只有几个字段，temperature 固定低值。"""

    def ir(self):
        return probe.scan("# 第一章 起点\n\n%s\n\n# 第二章 转折\n\n%s\n"
                          % (body(""), body("")))

    def test_accepts_known_kind(self):
        ir = probe.classify(self.ir(),
                            FakeLLM('{"kind":"methodology","unit_level":1,'
                                    '"reason":"论证链"}'),
                            CFG, _PARADIGMS)
        self.assertEqual(ir["kind"], "methodology")
        self.assertEqual(ir["unit_level"], 1)

    def test_unknown_kind_falls_back_to_auto(self):
        ir = probe.classify(self.ir(), FakeLLM('{"kind":"胡说","unit_level":1}'),
                            CFG, _PARADIGMS)
        self.assertEqual(ir["kind"], "auto")

    def test_model_failure_degrades_with_warning(self):
        # 探查是建议性的：模型不可用不该拦住人，但要说明按规则推断了
        ir = probe.classify(self.ir(), FakeLLM(boom=True), CFG, _PARADIGMS)
        self.assertTrue(any("类型判定未完成" in w for w in ir["warnings"]))

    def test_low_temperature(self):
        llm = FakeLLM('{"kind":"auto"}')
        probe.classify(self.ir(), llm, CFG, _PARADIGMS)
        self.assertLess(llm.calls[0]["temperature"], 0.5)

    def test_prompt_carries_the_genre_rules_for_level(self):
        # 判「按哪一级凝缩」的凭据是文体。只给类型的名字，等于让模型看着标题
        # 猜——标题越具体越像「一期讲一个」，层级就越判越细，凝缩次数跟着翻。
        llm = FakeLLM('{"kind":"methodology","unit_level":1}')
        probe.classify(self.ir(), llm, CFG, _PARADIGMS)
        prompt = llm.calls[0]["messages"][1]["content"]
        for key in ("凝缩单位", "切分依据", "整合依据", "重点判据", "推进方式",
                    "层级期望"):
            self.assertIn(key, prompt)
        # 代价要摆明：选哪一层，就是凝缩多少次
        self.assertIn("真正要凝缩的条数", prompt)
        # 结构按层级列出
        self.assertIn("H1（共", prompt)
        # 这一步不替人定数
        self.assertNotIn("正好排出", prompt)

    def test_records_which_model_decided(self):
        # 结论要留得下来处：同一份素材两次探查给出不同判定时，先要能分清是
        # 换了模型还是模型自身不稳。写了没记，等于这条线索不存在。
        ir = probe.classify(self.ir(),
                            FakeLLM('{"kind":"methodology","unit_level":1}'),
                            CFG, _PARADIGMS)
        self.assertEqual(ir["probe_model"], "fake")

    def test_frozen_ir_is_left_alone(self):
        ir = self.ir()
        ir["frozen"] = True
        llm = FakeLLM('{"kind":"methodology"}')
        probe.classify(ir, llm, CFG, _PARADIGMS)
        self.assertFalse(llm.calls)

    def test_unparseable_reply_degrades(self):
        ir = probe.classify(self.ir(), FakeLLM("我不想回答"), CFG, _PARADIGMS)
        self.assertTrue(any("按规则推断" in w for w in ir["warnings"]))


class TestScanProject(unittest.TestCase):
    def setUp(self):
        self.base = tempfile.mkdtemp(prefix="pm_probe_")
        self.item = P.create(self.base, "甲档", "mapped")
        self.pid = self.item["id"]

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def test_no_sources_raises(self):
        with self.assertRaises(probe.ProbeError):
            probe.scan_project(self.base, self.pid, CFG)

    def test_result_is_frozen_on_disk(self):
        # 探查与排图分开：排图失败（模型不可用、输出截断）不至于把探查赔进去
        S.add_source(self.base, self.pid, "书.md",
                     "# 第一章 起点\n\n%s\n" % body(""))
        probe.scan_project(self.base, self.pid, CFG)
        path = probe.probe_path(self.base, self.pid, "s1")
        self.assertTrue(os.path.exists(path))
        with open(path, "r", encoding="utf-8") as f:
            self.assertEqual(json.load(f)["segments"][0]["title"], "第一章 起点")

    def test_second_scan_reuses_the_file(self):
        S.add_source(self.base, self.pid, "书.md",
                     "# 第一章 起点\n\n%s\n" % body(""))
        probe.scan_project(self.base, self.pid, CFG)
        ir = probe.load(self.base, self.pid, "s1")
        self.assertIsNotNone(ir)

        # 把盘上结果改掉再扫一次：没重扫的话，改动会被读回来
        ir["segments"][0]["title"] = "改过的标题"
        probe.save(self.base, self.pid, "s1", ir)
        again = probe.scan_project(self.base, self.pid, CFG)
        self.assertEqual(again["sources"]["s1"]["segments"][0]["title"],
                         "改过的标题")

    def test_force_rescans(self):
        S.add_source(self.base, self.pid, "书.md",
                     "# 第一章 起点\n\n%s\n" % body(""))
        probe.scan_project(self.base, self.pid, CFG)
        ir = probe.load(self.base, self.pid, "s1")
        ir["segments"][0]["title"] = "改过的标题"
        probe.save(self.base, self.pid, "s1", ir)
        again = probe.scan_project(self.base, self.pid, CFG, force=True)
        self.assertEqual(again["sources"]["s1"]["segments"][0]["title"],
                         "第一章 起点")

    def test_summary_counts_body_not_misc(self):
        S.add_source(self.base, self.pid, "书.md",
                     "# 正文一\n\n%s\n\n# 附录 A\n\n%s\n" % (body(""), body("")))
        out = probe.scan_project(self.base, self.pid, CFG)
        ir = out["sources"]["s1"]
        self.assertLess(ir["body_chars"], ir["total_chars"])


class TestProjectCardReachesTheScan(unittest.TestCase):
    """项目定下的那张卡要一路走到探查。

    附属页是**扫标题那一步**判的，卡取不到就等于这一步没接。取卡走的是项目
    记录（立项时定的），不是类型判定——类型判定在扫结构之后，等它出来再回填，
    同一份素材的标题要判两遍，统计与告警还得跟着重算。
    """

    def setUp(self):
        self.base = tempfile.mkdtemp(prefix="pm_card_")
        self.pid = P.create(self.base, "甲档", "mapped", paradigm="paper")["id"]

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def misc_of(self, pid):
        S.add_source(self.base, pid, "书.md",
                     "# 正文一\n\n%s\n\n# 致谢\n\n%s\n" % (body(""), body("")))
        out = probe.scan_project(self.base, pid, CFG)
        return {s["title"]: s["misc"] for s in out["sources"]["s1"]["segments"]}

    def test_project_card_decides_the_appendix_list(self):
        self.assertEqual(self.misc_of(self.pid)["致谢"], "drop",
                         "项目定的是论文卡，致谢该按卡判为不播")

    def test_a_project_without_a_card_uses_the_shared_list(self):
        other = P.create(self.base, "乙档", "mapped")["id"]
        self.assertEqual(self.misc_of(other)["致谢"], "keep",
                         "项目没定类型，落通用名单：致谢是作者的话，要播")


class TestIngestMarks(unittest.TestCase):
    """docx 的样式证据要在抽取时就留下——取完只剩裸段落，就再也判不出层级。"""

    def test_docx_headings_become_anchors(self):
        try:
            import docx
        except ImportError:                                  # pragma: no cover
            self.skipTest("python-docx 不可用")
        base = tempfile.mkdtemp(prefix="pm_docx_")
        self.addCleanup(shutil.rmtree, base, True)
        path = os.path.join(base, "书.docx")
        d = docx.Document()
        d.add_heading("第一章 起点", level=1)
        d.add_paragraph("这一章讲事情的由来。" * 6)
        d.add_heading("第一节 名字的来历", level=2)
        d.add_paragraph("名字的来历是这样的。" * 6)
        d.save(path)

        text, meta = ingest.from_file(path)
        titles = [a["title"] for a in ingest.list_anchors(text, meta.get("marks"))]
        self.assertIn("第一章 起点", titles)
        self.assertIn("第一节 名字的来历", titles)
        # 层级来自样式，不是猜的
        lv = {a["title"]: a["level"] for a in
              ingest.list_anchors(text, meta.get("marks"))}
        self.assertLess(lv["第一章 起点"], lv["第一节 名字的来历"])

    def test_txt_numbering_becomes_anchors(self):
        base = tempfile.mkdtemp(prefix="pm_txt_")
        self.addCleanup(shutil.rmtree, base, True)
        path = os.path.join(base, "稿.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("第一章 起点\n\n%s\n\n第二章 转折\n\n%s\n" % (body(""), body("")))
        text, _meta = ingest.from_file(path)
        titles = [a["title"] for a in ingest.list_anchors(text)]
        self.assertIn("第一章 起点", titles)
        self.assertIn("第二章 转折", titles)

    def test_plain_txt_yields_no_anchor(self):
        # 纯文本什么结构都没有时，如实返回空，不编一套出来
        self.assertEqual(
            ingest.list_anchors(body("", 20)), [])

    def test_pdf_is_refused_with_a_way_out(self):
        base = tempfile.mkdtemp(prefix="pm_pdf_")
        self.addCleanup(shutil.rmtree, base, True)
        path = os.path.join(base, "书.pdf")
        with open(path, "wb") as f:
            f.write(b"%PDF-1.4\n")
        with self.assertRaises(ingest.IngestError) as ctx:
            ingest.from_file(path)
        msg = str(ctx.exception)
        self.assertIn("不支持 PDF", msg)
        self.assertIn("md", msg)


class TestClassifyBudget(unittest.TestCase):
    """类型判定是 L2 语义那一步，token 预算不能抠。

    推理型模型把思考过程也算进 max_tokens：预算给 800 时答案部分一个字也
    轮不上，分类永远退回规则推断——L2 这一层等于没跑，而日志里只留一句
    「模型不可用」，看着像模型的问题。
    """

    class _LLM(object):
        def __init__(self):
            self.calls = []

        def chat(self, messages, temperature=0.8, max_tokens=8192,
                 json_schema=None, **kw):
            self.calls.append({"max_tokens": max_tokens,
                               "temperature": temperature,
                               "schema": json_schema})
            return ('{"kind":"methodology","unit_level":1,"reason":"看得出"}',
                    {"model": "fake"})

    def _ir(self):
        return probe.scan("# 甲\n%s\n# 乙\n%s" % (body(""), body("")))

    def test_budget_comes_from_config_not_a_constant(self):
        ir = self._ir()
        llm = self._LLM()
        probe.classify(ir, llm, dict(CFG, **{"llm.max_tokens": 32768}),
                       _PARADIGMS)
        self.assertEqual(llm.calls[-1]["max_tokens"], 32768)

    def test_temperature_stays_low(self):
        # 分类要可复现：用排图那个温度，同一份素材两次探查给两套结果
        ir = self._ir()
        llm = self._LLM()
        probe.classify(ir, llm, dict(CFG, **{"llm.max_tokens": 32768}),
                       _PARADIGMS)
        self.assertEqual(llm.calls[-1]["temperature"], 0.1)


class TestBodyScope(unittest.TestCase):
    """字数的两把尺子，以及它们与取料的一致性。

    探查记的字数若不等于取料真正切走的量，排图就会按错的体量分期：有子节的
    章节记少了，体量算小了，期数算多了，而且那些章节还会被当成空壳挡在准入
    之外——它们从此不在任何一期里，而地图看上去是完整的。
    """

    TEXT = ("# 第一篇\n\n引子。\n\n## 第一章 甲\n\n" + "甲的内容。" * 30 +
            "\n\n## 第二章 乙\n\n" + "乙的内容。" * 30 + "\n")

    def test_self_contained_chars_match_the_slicing_rule(self):
        from podcast_maker import duration_model
        ir = probe.scan(self.TEXT)
        by = {s["title"]: s for s in ir["segments"]}
        body, _meta = ingest.slice_by_anchor(self.TEXT, "第一章 甲")
        want = round(duration_model.effective_chars(body)
                     - duration_model.effective_chars("第一章 甲"))
        self.assertEqual(by["第一章 甲"]["chars"], want)

    def test_parent_includes_its_subsections(self):
        ir = probe.scan(self.TEXT)
        by = {s["title"]: s for s in ir["segments"]}
        self.assertGreater(by["第一篇"]["chars"],
                           by["第一篇"]["own_chars"])
        self.assertGreaterEqual(by["第一篇"]["chars"],
                                by["第一章 甲"]["chars"] + by["第二章 乙"]["chars"])

    def test_own_chars_counts_only_the_direct_body(self):
        ir = probe.scan(self.TEXT)
        by = {s["title"]: s for s in ir["segments"]}
        # 「第一篇」的直属正文只有「引子。」这一句
        self.assertLess(by["第一篇"]["own_chars"], by["第一章 甲"]["own_chars"])

    def test_section_with_only_subsections_is_admissible(self):
        # 从前按直属正文判空壳，这种「标题下面直接是小节」的章全被挡掉
        ir = probe.scan(self.TEXT)
        ok, _bad = probe.admissible(ir)
        self.assertIn("第一篇", [s["title"] for s in ok])

    def test_units_are_one_level_not_two(self):
        # 父层与子层同时列出来，同一段正文要排两次期
        ir = probe.scan(self.TEXT)
        ir["unit_level"] = 2
        got = [s["title"] for s in probe.units(ir)]
        self.assertEqual(got, ["第一章 甲", "第二章 乙"])

    def test_body_chars_excludes_appendix_subsections(self):
        # 「默认不播」的字样只写在附录自己的标题上，它的小节标题不带
        text = ("# 正文\n\n" + "正文内容。" * 30 +
                "\n\n# 附录 A\n\n## 甲表\n\n" + "表的内容。" * 30 + "\n")
        ir = probe.scan(text)
        self.assertLess(ir["body_chars"], ir["total_chars"])
        self.assertLess(ir["body_chars"], probe.scan(text)["total_chars"])


class TestPickUnits(unittest.TestCase):
    """取料单元：切分单位那一层；装不下就按下一层展开。"""

    def test_oversized_unit_is_split_by_its_own_structure(self):
        book = ("# 大部头\n\n" + "开头。" * 10 + "\n\n" +
                "".join("## 第%d节\n\n" % i + "内容。" * 200 + "\n\n"
                        for i in range(1, 5)))
        ir = probe.scan(book)
        ir["unit_level"] = 1
        got = probe.pick_units(ir, per_ep=300)
        self.assertEqual([s["title"] for s in got],
                         ["第1节", "第2节", "第3节", "第4节"])

    def test_unit_that_fits_is_left_whole(self):
        ir = probe.scan("# 甲\n\n短。\n\n# 乙\n\n也短。\n")
        got = probe.pick_units(ir, per_ep=100000)
        self.assertEqual([s["title"] for s in got], ["甲", "乙"])

    def test_unsplittable_unit_is_flagged_not_dropped(self):
        # 已经是最后一层，拆无可拆：整块交出去并把「超时长」标出来，
        # 由排图阶段告警，不是悄悄丢掉这一节
        ir = probe.scan("# 巨章\n\n" + "内容。" * 400 + "\n")
        ir["unit_level"] = 1
        got = probe.pick_units(ir, per_ep=100)
        self.assertEqual([s["title"] for s in got], ["巨章"])
        self.assertTrue(got[0]["oversize"])

    def test_appendix_and_its_subsections_stay_out(self):
        text = ("# 正文\n\n" + "正文内容。" * 30 +
                "\n\n# 附录 A\n\n## 甲表\n\n" + "表的内容。" * 30 + "\n")
        ir = probe.scan(text)
        ir["unit_level"] = 2
        got = [s["title"] for s in probe.pick_units(ir, per_ep=100000)]
        self.assertNotIn("甲表", got)
        self.assertNotIn("附录 A", got)


class TestDuplicateTitles(unittest.TestCase):
    """同名标题：取料要按行号定到具体那一条。

    各章都有一节「小结」是常态。只按标题取，取到的是靠前那一条，而地图上写的
    是同一串字——错了也看不出来。
    """

    TEXT = ("# 甲章\n\n第一节的内容。\n\n# 小结\n\n甲章的小结。\n\n"
            "# 乙章\n\n第二节的内容。\n\n# 小结\n\n乙章的小结。\n")

    def test_scan_warns_about_duplicates(self):
        ir = probe.scan(self.TEXT)
        self.assertTrue(any("同名标题" in w for w in ir["warnings"]),
                        ir["warnings"])

    def test_line_picks_the_second_one(self):
        ir = probe.scan(self.TEXT)
        second = [s for s in ir["segments"] if s["title"] == "小结"][1]
        body, _meta = ingest.slice_by_anchor(self.TEXT, "小结",
                                             line=second["line"])
        self.assertIn("乙章的小结", body)

    def test_without_line_it_falls_back_to_the_first(self):
        body, _meta = ingest.slice_by_anchor(self.TEXT, "小结")
        self.assertIn("甲章的小结", body)

    def test_recorded_chars_are_for_the_segment_not_the_first_match(self):
        from podcast_maker import duration_model
        ir = probe.scan(self.TEXT)
        second = [s for s in ir["segments"] if s["title"] == "小结"][1]
        body, _meta = ingest.slice_by_anchor(self.TEXT, "小结",
                                             line=second["line"])
        want = round(duration_model.effective_chars(body)
                     - duration_model.effective_chars("小结"))
        self.assertEqual(second["chars"], want)


class TestCondense(unittest.TestCase):
    """逐单元凝缩：一次调用只看一节，结果写回单元并落盘。"""

    def setUp(self):
        self.base = tempfile.mkdtemp(prefix="pm_cond_")
        self.item = P.create(self.base, "甲档", "mapped")
        self.pid = self.item["id"]
        S.add_source(self.base, self.pid, "书.md", BOOK)

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def units(self):
        ir = probe.load(self.base, self.pid, "s1") or probe.scan(BOOK)
        return [(s["title"], s) for s in probe.pick_units(ir, 100000)]

    def test_one_call_per_unit(self):
        llm = FakeLLM('{"gist":"这一节的主旨","points":["重点"],"concepts":["概念"]}')
        ir = probe.scan(BOOK)
        probe.save(self.base, self.pid, "s1", ir)
        done, skipped = probe.condense_units(
            self.base, self.pid,
            [("s1", s) for s in probe.pick_units(ir, 100000)],
            llm, CFG, _PARADIGMS["methodology"])
        self.assertEqual(done, 3)
        self.assertEqual(skipped, 0)
        self.assertEqual(len(llm.calls), 3)
        for c in llm.calls:
            self.assertEqual(c["schema"], probe.CONDENSE_SCHEMA)
            self.assertLess(c["temperature"], 0.5)

    def test_result_is_written_back_and_reused(self):
        llm = FakeLLM('{"gist":"主旨甲","points":["重点"],"concepts":[]}')
        ir = probe.scan(BOOK)
        units = [("s1", s) for s in probe.pick_units(ir, 100000)]
        probe.condense_units(self.base, self.pid, units, llm, CFG,
                             _PARADIGMS["methodology"])
        self.assertEqual(units[0][1]["gist"], "主旨甲")
        probe.save(self.base, self.pid, "s1", ir)
        again = probe.load(self.base, self.pid, "s1")
        self.assertEqual(again["segments"][0]["gist"], "主旨甲")
        # 第二次不再调模型：重排一次图不该把 N 次调用再烧一遍
        llm2 = FakeLLM('{"gist":"改过的主旨"}')
        done, skipped = probe.condense_units(
            self.base, self.pid,
            [("s1", s) for s in probe.pick_units(again, 100000)],
            llm2, CFG, _PARADIGMS["methodology"])
        self.assertEqual(llm2.calls, [])
        self.assertEqual((done, skipped), (0, 3))

    def test_force_recondenses(self):
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        units = [("s1", s) for s in probe.pick_units(ir, 100000)]
        probe.condense_units(self.base, self.pid, units, llm, CFG,
                             _PARADIGMS["methodology"])
        llm2 = FakeLLM('{"gist":"改过的主旨"}')
        done, _sk = probe.condense_units(
            self.base, self.pid,
            [("s1", s) for s in probe.pick_units(ir, 100000)], llm2, CFG,
            _PARADIGMS["methodology"], force=True)
        self.assertEqual(done, 3)

    def test_failure_is_fatal_not_silent(self):
        # 凝缩结果是排图分组的依据。拿不到它还要往下走，排出来的图是按字数
        # 硬切的——假功能比报错坏得多。
        ir = probe.scan(BOOK)
        units = [("s1", s) for s in probe.pick_units(ir, 100000)]
        with self.assertRaises(probe.ProbeError):
            probe.condense_units(self.base, self.pid, units,
                                 FakeLLM("我不想凝缩"), CFG,
                                 _PARADIGMS["methodology"])

    def test_prompt_carries_the_condense_rule_not_the_merge_rule(self):
        # 凝缩与排图是两件事，判据方向相反：凝缩要「不丢」，排图要「选」。
        # 把排图那条（focus）喂给凝缩，模型会先按「挑重点」办事，该留的
        # 细节区分被当成次要信息丢掉——而排图那一步只看凝缩、看不见原文。
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", probe.pick_units(ir, 100000)[0])],
                             llm, CFG, _PARADIGMS["methodology"])
        prompt = llm.calls[0]["messages"][1]["content"]
        card = _PARADIGMS["methodology"]
        self.assertIn(card["condense"][:12], prompt)
        self.assertNotIn(card["focus"][:12], prompt)

    def test_prompt_keeps_the_details_that_decide_merging(self):
        # 「说明了 X 的重要性」这种空话对排图没用：看不出两节该合还是该拆。
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", probe.pick_units(ir, 100000)[0])],
                             llm, CFG, _PARADIGMS["methodology"])
        prompt = llm.calls[0]["messages"][1]["content"]
        for want in ("反对什么", "细节区分", "有几条主线"):
            self.assertIn(want, prompt)

    def test_prompt_gives_a_worked_example(self):
        # 只有字段说明不够：凝缩是**每个单元一次调用**，缺一条能照抄的示例，
        # 格式上的偏差会被乘上单元数，而这一步失败是致命的。
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", probe.pick_units(ir, 100000)[0])],
                             llm, CFG, _PARADIGMS["methodology"])
        prompt = llm.calls[0]["messages"][1]["content"]
        self.assertIn("示例", prompt)
        obj = llm_mod.extract_json(prompt[prompt.index("示例"):])
        self.assertIsNotNone(obj, "示例本身要是一段能解析的 JSON")
        self.assertIn("gist", obj)
        self.assertIn("points", obj)

    def test_condensation_does_not_decide_thinking_for_the_user(self):
        # 用推理型还是普通模型由使用者自己定。这里只管把要求写清楚、把结果
        # 接住——不传「别思考」这类字段，避免替使用者做决定。
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", s) for s in probe.pick_units(ir, 100000)],
                             llm, CFG, _PARADIGMS["methodology"])
        for c in llm.calls:
            self.assertIsNone(c.get("reasoning_effort"))

    def test_body_comes_from_the_slicing_rule(self):
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", probe.pick_units(ir, 100000)[0])],
                             llm, CFG, _PARADIGMS["methodology"])
        prompt = llm.calls[0]["messages"][1]["content"]
        self.assertIn("这一章讲事情的由来", prompt)

    def test_prompt_sets_no_length_cap(self):
        # 「每条 40 字内」把逻辑链压成了口号，细节区分一丢，两节看着就像在讲
        # 同一件事——而排图正是靠这个判该合还是该拆。
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", probe.pick_units(ir, 100000)[0])],
                             llm, CFG, _PARADIGMS["methodology"])
        prompt = llm.calls[0]["messages"][1]["content"]
        self.assertNotIn("40 字", prompt)
        self.assertIn("不设条数上限", prompt)

    def test_every_paradigm_rule_survives_the_prompt_template(self):
        # 提示词整段是用 % 拼的：卡里出现一个裸百分号（比如写「超标 40%」），
        # 运行时就会炸，而报错指向的是格式化、不是那张卡——追起来很远。
        for key, card in _PARADIGMS.items():
            probe._condense_prompt("标题", card.get("condense") or "", "正文")

    def test_result_carries_the_version_it_was_made_with(self):
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        units = [("s1", s) for s in probe.pick_units(ir, 100000)]
        probe.condense_units(self.base, self.pid, units, llm, CFG,
                             _PARADIGMS["methodology"])
        self.assertEqual(units[0][1]["condense_version"], probe.CONDENSE_VERSION)

    def test_result_from_an_older_version_is_redone(self):
        # 凝缩的取舍标准写在提示词里。提示词换过，旧结果就是另一套标准下的
        # 产物；留着它，一张图上会混着两套标准，从界面上看不出来。
        llm = FakeLLM('{"gist":"新主旨"}')
        ir = probe.scan(BOOK)
        units = [("s1", s) for s in probe.pick_units(ir, 100000)]
        for _sid, s in units:
            s["gist"] = "上一版留下的主旨"
            s["condense_version"] = probe.CONDENSE_VERSION
        units[0][1]["condense_version"] = "0"
        done, skipped = probe.condense_units(self.base, self.pid, units, llm,
                                             CFG, _PARADIGMS["methodology"])
        self.assertEqual(done, 1)
        self.assertEqual(skipped, 2)
        self.assertEqual(units[0][1]["gist"], "新主旨")
        self.assertEqual(units[1][1]["gist"], "上一版留下的主旨")

    def test_stamp_records_which_model_and_version_made_the_gist(self):
        llm = FakeLLM('{"gist":"主旨甲"}')
        ir = probe.scan(BOOK)
        probe.condense_units(self.base, self.pid,
                             [("s1", s) for s in probe.pick_units(ir, 100000)],
                             llm, CFG, _PARADIGMS["methodology"])
        probe.stamp_condense(ir)
        self.assertEqual(ir["condense_version"], probe.CONDENSE_VERSION)
        self.assertEqual(ir["condense_model"], "fake")

    def test_stamp_leaves_untouched_ir_empty(self):
        ir = probe.scan(BOOK)
        probe.stamp_condense(ir)
        self.assertEqual(ir["condense_version"], "")
        self.assertEqual(ir["condense_model"], "")


class TestStaleIr(unittest.TestCase):
    """旧版本的探查结果按「没探查过」处理。

    旧版只有一把字数尺子（直属正文），自含体量与取料对不上：拿它排图会把有
    子节的章节算小、把附录的子节算进正文。宁可重扫一次，也不把两种口径混在
    一张图上。
    """

    def setUp(self):
        self.base = tempfile.mkdtemp(prefix="pm_stale_")
        self.item = P.create(self.base, "甲档", "mapped")
        self.pid = self.item["id"]
        S.add_source(self.base, self.pid, "书.md", BOOK)

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def test_old_version_is_treated_as_unscanned(self):
        ir = probe.scan(BOOK)
        ir["version"] = 1
        probe.save(self.base, self.pid, "s1", ir)
        self.assertIsNone(probe.load(self.base, self.pid, "s1"))

    def test_scan_project_rescans_old_results(self):
        ir = probe.scan(BOOK)
        ir["version"] = 1
        probe.save(self.base, self.pid, "s1", ir)
        out = probe.scan_project(self.base, self.pid, CFG)
        self.assertEqual(out["sources"]["s1"]["version"], probe.IR_VERSION)


from podcast_maker import paradigms as _PG                        # noqa: E402

#: 分类函数收的是范式字典（键 → 卡），不是范式模块
_PARADIGMS = _PG.PARADIGMS


if __name__ == "__main__":
    unittest.main()
